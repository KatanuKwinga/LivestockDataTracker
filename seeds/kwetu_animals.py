"""Bit 3: load Kwetu's animals from the workbook into one farm."""
from openpyxl import load_workbook

from app import create_app
from app.extensions import db
from app.models import Livestock, User

EMAIL = "katanukavilakwinga@gmail.com"
WORKBOOK = "seeds/data/Kwetu_Animals_and_Records.xlsx"
MARK = "[kwetu]"

SPECIES = {"Cattle": "CATTLE", "Goat": "GOAT", "Sheep": "SHEEP"}
SEX = {"Male": "MALE", "Female": "FEMALE"}

app = create_app()

with app.app_context():
    farm_id = User.query.filter_by(email=EMAIL).first().farmer.farmer_id

    sheet = load_workbook(WORKBOOK, read_only=True)["Animals"]
    added = 0
    for row in sheet.iter_rows(min_row=2, values_only=True):
        tag, species, sex, breed, born = row[0], row[1], row[2], row[3], row[4]
        if tag is None:
            continue
        animal = Livestock(
            farmer_id=farm_id,
            tag_number=tag,
            species=SPECIES[species],
            gender=SEX[sex],
            breed=breed,
            date_of_birth=born.date(),
            description=f"{MARK} {row[8]}",
        )
        db.session.add(animal)
        added += 1

    db.session.commit()
    print("added", added, "animals")